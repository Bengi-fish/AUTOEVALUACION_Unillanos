#!/usr/bin/env python3
"""
Convierte un plan de mejoramiento en formato FO-GCL-20 (Excel de Unillanos) a CSV
listos para importar en Drupal (Migrate API) según el MER v2.

Uso:
    python3 datos/herramientas/plan_excel_a_csv.py RUTA_AL_EXCEL.xlsx datos/importacion/<slug-programa>/

Genera en la carpeta de salida:
    plan.csv                 encabezado del plan (PLAN_MEJORAMIENTO)
    hallazgos_plan.csv       oportunidades de mejora del plan (HALLAZGO con destino = plan de mejoramiento)
    metas.csv                metas (META) con su relación N:M a hallazgos y responsables
    indicadores.csv          indicadores (INDICADOR)
    programacion_anual.csv   valor programado por indicador y año (PROGRAMACION_ANUAL)
    seguimientos.csv         reportes semestrales (SEGUIMIENTO), una fila por meta y hoja semestral
    responsables.csv         lista única de responsables encontrados (RESPONSABLE)
    listas_facultades.csv, listas_programas.csv, listas_proyectos.csv   (hoja oculta "Listas")

Requiere: pip install openpyxl
Los identificadores (id) son claves de importación estables dentro de este plan (p. ej. M01, M01-I1),
no los IDs internos de Drupal.
"""
import csv, os, re, sys
from datetime import date, datetime

try:
    import openpyxl
except ImportError:
    sys.exit("Falta openpyxl: pip install openpyxl")

HOJA_PLAN = "Plan mejoramiento"
FILA_DATOS = 10          # primera fila de metas (las filas 8-9 son encabezados)
COL = dict(factor=1, origen=2, descripcion=3, proyecto=4, peso=5, tipo_meta=6, meta=7, indicador=8,
           actividades=9, recursos=10, linea_base=11, base_valor=12, anio1=13, total=20, responsable=21)
ANIO_INICIAL = 2024      # "Año 1" del formato 2024-2 – 2030


def limpio(v):
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"[ \t]+", " ", str(v)).strip()


def una_linea(v):
    return re.sub(r"\s+", " ", limpio(v))


class Hoja:
    """Lee celdas resolviendo celdas combinadas (devuelve el valor de la esquina superior izquierda)."""

    def __init__(self, ws):
        self.ws = ws
        self.ancla = {}
        for rng in ws.merged_cells.ranges:
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    self.ancla[(r, c)] = (rng.min_row, rng.min_col)

    def celda(self, r, c):
        return self.ancla.get((r, c), (r, c))

    def valor(self, r, c):
        rr, cc = self.celda(r, c)
        return self.ws.cell(rr, cc).value


def escribir(carpeta, nombre, filas, campos):
    with open(os.path.join(carpeta, nombre), "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=campos)
        w.writeheader()
        for fila in filas:
            w.writerow({k: fila.get(k, "") for k in campos})


def buscar_etiqueta(h, etiqueta, filas=range(1, 10), cols=range(1, 22)):
    """Valor de la celda a la derecha de una etiqueta del encabezado (p. ej. 'Programa:')."""
    for r in filas:
        for c in cols:
            if una_linea(h.valor(r, c)).lower().startswith(etiqueta.lower()):
                for c2 in range(c + 1, c + 4):
                    v = una_linea(h.valor(r, c2))
                    if v and h.celda(r, c2) != h.celda(r, c):
                        return v
    return ""


def clave_meta(texto):
    """Clave para emparejar la meta del plan con la de las hojas semestrales."""
    return una_linea(texto).lower()[:40]


def main(ruta, salida):
    os.makedirs(salida, exist_ok=True)
    wb = openpyxl.load_workbook(ruta, data_only=True)
    h = Hoja(wb[HOJA_PLAN])

    # ---------- encabezado ----------
    plan = {
        "id": "P01",
        "tipo": buscar_etiqueta(h, "Tipo de plan"),
        "programa": buscar_etiqueta(h, "Programa:"),
        "facultad": buscar_etiqueta(h, "Facultad:"),
        "nivel": buscar_etiqueta(h, "Nivel de formación"),
        "proceso": buscar_etiqueta(h, "Proceso:"),
        "alta_calidad": buscar_etiqueta(h, "Alta calidad"),
        "fecha_suscripcion": buscar_etiqueta(h, "Fecha de suscripción"),
        "director": buscar_etiqueta(h, "Director de programa"),
        "codigo_formato": "", "version_formato": "", "periodo": "",
    }
    for c in range(1, 22):
        v = una_linea(h.valor(3, c))
        if v.startswith("Código:"):
            plan["codigo_formato"] = v.split(":", 1)[1].strip()
        if v.startswith("Versión:"):
            plan["version_formato"] = v.split(":", 1)[1].strip()
    m = re.search(r"(\d{4}-\d)\s*-\s*(\d{4})", una_linea(h.valor(2, 2)))
    if m:
        plan["periodo"] = f"{m.group(1)} – {m.group(2)}"

    # ---------- filas de metas / indicadores ----------
    metas, hallazgos, indicadores, programacion = {}, {}, [], []
    responsables = {}
    ultima = h.ws.max_row
    for r in range(FILA_DATOS, ultima + 1):
        if una_linea(h.valor(r, 1)).lower().startswith("elaborado"):
            break
        indicador = limpio(h.valor(r, COL["indicador"]))
        meta_txt = limpio(h.valor(r, COL["meta"]))
        if not indicador and not meta_txt:
            continue

        # Meta (se identifica por su celda combinada)
        mk = h.celda(r, COL["meta"])
        if mk not in metas:
            n = len(metas) + 1
            cod = re.match(r"Meta\s*([^\n]+?)\.?\s*\n", meta_txt)
            resp = [una_linea(x).rstrip(".").strip() for x in limpio(h.valor(r, COL["responsable"])).split("\n") if una_linea(x)]
            for x in resp:
                responsables.setdefault(x.lower(), x)
            metas[mk] = {
                "id": f"M{n:02d}", "plan_id": plan["id"], "numero": n,
                "codigo": cod.group(1).strip() if cod else str(n),
                "descripcion": re.sub(r"^Meta\s*[^\n]*\n", "", meta_txt).strip() or meta_txt,
                "factor_texto": una_linea(h.valor(r, COL["factor"])),
                "proyecto": una_linea(h.valor(r, COL["proyecto"])),
                "tipo_meta": una_linea(h.valor(r, COL["tipo_meta"])),
                "peso": limpio(h.valor(r, COL["peso"])),
                "actividades": limpio(h.valor(r, COL["actividades"])),
                "recursos": una_linea(h.valor(r, COL["recursos"])),
                "responsables": "|".join(resp),
                "hallazgos": [],
                "_n_ind": 0,
            }
        meta = metas[mk]

        # Oportunidad de mejora = HALLAZGO (N:M con meta)
        hk = h.celda(r, COL["descripcion"])
        desc = una_linea(h.valor(r, COL["descripcion"]))
        if desc:
            if hk not in hallazgos:
                hallazgos[hk] = {
                    "id": f"H{len(hallazgos) + 1:02d}", "tipo": "aspecto_por_mejorar",
                    "descripcion": desc,
                    "origen": una_linea(h.valor(r, COL["origen"])),
                    "factor_texto": una_linea(h.valor(r, COL["factor"])),
                    "destino": "plan_mejoramiento",
                }
            if hallazgos[hk]["id"] not in meta["hallazgos"]:
                meta["hallazgos"].append(hallazgos[hk]["id"])

        # Indicador + programación anual
        if indicador:
            meta["_n_ind"] += 1
            iid = f"{meta['id']}-I{meta['_n_ind']}"
            indicadores.append({
                "id": iid, "meta_id": meta["id"], "nombre": una_linea(indicador),
                "linea_base": una_linea(h.valor(r, COL["linea_base"])),
                "valor_base": limpio(h.valor(r, COL["base_valor"])),
                "total": limpio(h.ws.cell(r, COL["total"]).value),
            })
            for k in range(7):
                v = h.ws.cell(r, COL["anio1"] + k).value
                if v not in (None, ""):
                    programacion.append({"indicador_id": iid, "anio": ANIO_INICIAL + k, "programado": limpio(v), "logrado": ""})

    # ---------- hojas semestrales = SEGUIMIENTO ----------
    por_clave = {clave_meta(m["descripcion"]): m for m in metas.values()}
    por_codigo = {m["codigo"].lower(): m for m in metas.values()}
    seguimientos = []
    for ws in wb.worksheets:
        if not re.fullmatch(r"\d{4}-[12]", ws.title):
            continue
        hs = Hoja(ws)
        vistos = set()
        for r in range(FILA_DATOS, ws.max_row + 1):
            mt = limpio(hs.valor(r, 2))
            if not mt or hs.celda(r, 2) in vistos:
                continue
            vistos.add(hs.celda(r, 2))
            cod = re.match(r"Meta\s*([^\n]+?)\.?\s*\n", mt)
            meta = por_codigo.get(cod.group(1).strip().lower()) if cod else None
            if not meta:
                meta = por_clave.get(clave_meta(re.sub(r"^Meta\s*[^\n]*\n", "", mt)))
            if not meta:
                continue
            fila = {
                "meta_id": meta["id"], "periodo": ws.title,
                "avance": limpio(hs.valor(r, 3)),
                "actividades": limpio(hs.valor(r, 4)),
                "medios_verificacion": limpio(hs.valor(r, 5)),
                "verificacion": limpio(hs.valor(r, 6)),
                "observaciones": limpio(hs.valor(r, 7)),
            }
            if any(fila[k] for k in ("avance", "actividades", "verificacion", "observaciones")):
                seguimientos.append(fila)

    # ---------- hoja oculta "Listas" ----------
    fac, progs, proys = [], [], []
    if "Listas" in wb.sheetnames:
        hl = Hoja(wb["Listas"])
        siglas, niveles = {}, {}
        for c in range(1, 26):
            s = una_linea(hl.valor(2, c))
            if re.fullmatch(r"FC[A-Z]+", s):
                siglas[c] = s
        sigla_col = {}
        actual = None
        for c in range(4, 20):
            if c in siglas:
                actual = siglas[c]
            n = una_linea(hl.valor(3, c))
            if actual and n:
                sigla_col[c] = (actual, n)
        nombres_fac = [una_linea(hl.valor(r, 2)) for r in range(4, 12) if una_linea(hl.valor(r, 2))]
        orden_siglas = list(dict.fromkeys(s for s, _ in sigla_col.values()))
        for i, nombre in enumerate(nombres_fac):
            fac.append({"sigla": orden_siglas[i] if i < len(orden_siglas) else "", "nombre": nombre})
        for c, (sig, niv) in sigla_col.items():
            for r in range(4, hl.ws.max_row + 1):
                v = una_linea(hl.ws.cell(r, c).value)
                if v:
                    progs.append({"nombre": v, "facultad_sigla": sig, "nivel": niv})
        for r in range(4, hl.ws.max_row + 1):
            v = una_linea(hl.ws.cell(r, 24).value)
            if v:
                proys.append({"nombre": v})

    # ---------- escribir ----------
    escribir(salida, "plan.csv", [plan], list(plan.keys()))
    escribir(salida, "hallazgos_plan.csv", hallazgos.values(), ["id", "tipo", "descripcion", "origen", "factor_texto", "destino"])
    for m in metas.values():
        m["hallazgos"] = "|".join(m["hallazgos"])
    escribir(salida, "metas.csv", metas.values(), ["id", "plan_id", "numero", "codigo", "descripcion", "factor_texto", "proyecto",
                                                    "tipo_meta", "peso", "actividades", "recursos", "responsables", "hallazgos"])
    escribir(salida, "indicadores.csv", indicadores, ["id", "meta_id", "nombre", "linea_base", "valor_base", "total"])
    escribir(salida, "programacion_anual.csv", programacion, ["indicador_id", "anio", "programado", "logrado"])
    escribir(salida, "seguimientos.csv", seguimientos, ["meta_id", "periodo", "avance", "actividades", "medios_verificacion",
                                                        "verificacion", "observaciones"])
    escribir(salida, "responsables.csv", [{"nombre": v} for v in sorted(responsables.values())], ["nombre"])
    if fac:
        escribir(salida, "listas_facultades.csv", fac, ["sigla", "nombre"])
        escribir(salida, "listas_programas.csv", progs, ["nombre", "facultad_sigla", "nivel"])
        escribir(salida, "listas_proyectos.csv", proys, ["nombre"])

    print(f"Plan: {plan['programa'] or plan['tipo']} · {len(metas)} metas · {len(hallazgos)} oportunidades de mejora · "
          f"{len(indicadores)} indicadores · {len(programacion)} valores anuales · {len(seguimientos)} seguimientos · "
          f"{len(responsables)} responsables")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])
